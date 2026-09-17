"""spec-v1.9.0 T8 (docs/spec/spec-v1.9.0.md Sec.10, REQ-V190-EVAL-01..04):
gate 7, the retrieval evaluation. Indexes the four-document eval corpus for
a synthetic user through `documents.index_document` (the real production
pipeline), then measures `vector`/`hybrid`/`hybrid+rerank` retrieval against
a frozen 12-question set, plus an advisory conversation-aware smoke test.

`devtools/` is never imported by the bot (AGENTS.md); this module talks to
the live LM Studio box exactly as `bot.py --selftest-live` does, through
`Config`/`EmbeddingsClient`/`build_llm_client` -- no shortcuts, no mocking of
the retrieval path itself. `run()` is the whole script's logic, parameterised
over its I/O boundaries (`conn`, `embedder`, `llm`, `cfg`, corpus/questions
paths) so `tests/test_v190_eval.py` (`T-V190-EVAL-02`) can drive it offline
with `FakeEmbedder`/`FakeLLM`; `main()` only wires the real ones.
"""

from __future__ import annotations

import dataclasses
import hashlib
import io
import json
import logging
import re
import sys
import tempfile
import time
from collections.abc import Callable
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    # `rag_eval.py` is invoked as a script (`uv run --locked python
    # devtools/rag_eval.py`, AGENTS.md's gate 7), so the project root is not
    # on `sys.path` by default -- matches `devtools/bench.py`'s own fix.
    sys.path.insert(0, str(REPO_ROOT))

import docx  # noqa: E402
import httpx  # noqa: E402

import agent  # noqa: E402
import config  # noqa: E402
import documents  # noqa: E402
import rag  # noqa: E402
import storage  # noqa: E402
from config import Config, load_config  # noqa: E402
from devtools.pdf_fixture import write_pdf  # noqa: E402
from llm import build_llm_client  # noqa: E402
from llm.base import LLMError  # noqa: E402
from llm.embeddings import EmbeddingError, EmbeddingsClient, EmbeddingTimeoutError  # noqa: E402

CORPUS_DIR = REPO_ROOT / "evals" / "rag" / "corpus"
QUESTIONS_PATH = REPO_ROOT / "evals" / "rag" / "questions.json"

# REQ-V190-EVAL-03: the synthetic tenant this script indexes and queries
# under -- never a real Telegram user id.
EVAL_USER_ID = -1

RECALL_FLOOR = 0.8
MODES = ("vector", "hybrid", "hybrid+rerank")

# The page-separator convention for `security_guidelines.pdf.txt`: a line
# whose content is exactly one form-feed character (never stripped by a
# whitespace-collapsing split, since blank lines are already "").
_PAGE_BREAK = "\x0c"

_WHITESPACE_RE = re.compile(r"\s+")
_TABLE_ROW_SEP_RE = re.compile(r"^-+$")

# The corpus's own filename contract (REQ-V190-EVAL-01): the four committed
# source files, and the filename `documents.index_document` sees for each --
# the DOCX/PDF pair is rendered from its committed text source at index time,
# never committed as binary.
_CORPUS_SOURCES: tuple[tuple[str, str, str], ...] = (
    ("vacation_policy.md", "vacation_policy.md", "text"),
    ("onboarding.txt", "onboarding.txt", "text"),
    ("expenses.docx.md", "expenses.docx", "docx"),
    ("security_guidelines.pdf.txt", "security_guidelines.pdf", "pdf"),
)

# TOOL-06's advisory smoke test: the brief's literal follow-up text. Turn
# 1's own question is the run's first answerable item (`run()` supplies
# it), never hardcoded here, so the smoke test stays corpus-agnostic.
_SMOKE_FOLLOWUP = "а в неделях?"

# v1.9.4 T2 (docs/spec/task-briefs/v194-T2.md): turn 2's own verdict cannot
# fail for a real reason today -- every run since v1.9.3 T2's fix answers
# "28 дней -- это 4 недели" from turn 1's own context without a
# `search_documents` call, a *correct* answer the token-overlap check
# reports as `fail`, indistinguishable from a real REQ-V190-TOOL-06
# regression. Turn 3 is a paraphrase of `evals/rag/questions.json`'s second
# answerable item (gold `vacation_policy.md`) phrased as a follow-up that
# cannot be answered from turns 1-2's own context, so its verdict --
# whether the retrieved passages actually include the gold source -- is
# meaningful pass/fail, never merely advisory-by-construction.
_SMOKE_FOLLOWUP_2 = "а сколько из них можно перенести на следующий год?"
# evals/rag/questions.json item 1's expected_source -- unlike `question`
# above, this is a real-corpus constant, not threaded through from `run()`
# (`answerable[1]` would resolve to a different filename against the
# synthetic corpus's own _SYNTH_QUESTIONS order, flipping the gold source
# the offline tests exercise). `tests/test_v190_eval.py`'s
# `_write_synth_corpus` deliberately reuses this same filename among its
# four synthetic documents so this constant, not a corpus-specific one, is
# what those tests exercise too -- swapping the real eval corpus's second
# answerable item's source would need this constant updated by hand.
_SMOKE_GOLD_SOURCE_2 = "vacation_policy.md"
# evals/rag/questions.json item 1's expected_evidence -- like _SMOKE_GOLD_SOURCE_2,
# a real-corpus constant, not threaded through from run(). Tightens turn 3's
# verdict to match first_hit's own matcher (filename AND evidence, not
# filename alone): the same file passing on filename alone would be true
# for ANY vacation-related search (already proven by turn 1's own hit),
# never demonstrating that the *follow-up*'s specific answer -- the
# transfer-limit rule -- was actually retrieved.
_SMOKE_GOLD_EVIDENCE_2 = "не более 10 дней"


# ---------------------------------------------------------------------------
# Normalisation and hashing shared by the offline shape test (EVAL-01/-04)
# and the live scoring (EVAL-03).
# ---------------------------------------------------------------------------


def normalize(text: str) -> str:
    """Case-insensitive, whitespace-normalised form used for every
    evidence/hit comparison in this module."""
    return _WHITESPACE_RE.sub(" ", text).strip().casefold()


def contains_evidence(text: str, evidence: str) -> bool:
    return normalize(evidence) in normalize(text)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def freeze(corpus_dir: Path, questions_path: Path) -> dict[str, str]:
    """REQ-V190-EVAL-01/-02's freeze: sha256 of every corpus file and of
    `questions.json`, computed once, before the first live call."""
    hashes = {
        path.name: sha256_file(path) for path in sorted(corpus_dir.glob("*")) if path.is_file()
    }
    hashes[questions_path.name] = sha256_file(questions_path)
    return hashes


# ---------------------------------------------------------------------------
# Corpus construction: the two committed text sources are indexed verbatim;
# the DOCX/PDF pair is rendered in memory from its committed text source.
# ---------------------------------------------------------------------------


def _parse_md_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [cell.strip() for cell in cells]


def _is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(_TABLE_ROW_SEP_RE.match(cell) for cell in cells)


def markdown_to_docx_bytes(source: str) -> bytes:
    """A minimal Markdown-paragraphs-and-pipe-tables renderer, just rich
    enough for `expenses.docx.md`: a block (separated by a blank line) whose
    every non-empty line starts with `|` becomes a table (its `|---|---|`
    separator row dropped); every other block becomes one paragraph, its
    internal line breaks collapsed to spaces so `documents._extract_docx`'s
    per-paragraph text stays one line."""
    document = docx.Document()
    blocks = re.split(r"\n\s*\n", source.strip())
    for block in blocks:
        lines = [line for line in block.split("\n") if line.strip()]
        if not lines:
            continue
        if all(line.strip().startswith("|") for line in lines):
            rows = [_parse_md_table_row(line) for line in lines]
            rows = [row for row in rows if not _is_separator_row(row)]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row_cells in enumerate(rows):
                for c, value in enumerate(row_cells):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(" ".join(line.strip() for line in lines))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def split_pdf_pages(source: str) -> list[str]:
    """Splits on a line that is *exactly* one form-feed character -- not
    `.strip()`-equal, since a blank line (`""`) would also strip to `""` and
    `str.strip()` treats `\\x0c` itself as whitespace it would remove."""
    lines = source.split("\n")
    pages: list[str] = []
    current: list[str] = []
    for line in lines:
        if line == _PAGE_BREAK:
            pages.append("\n".join(current).strip())
            current = []
        else:
            current.append(line)
    pages.append("\n".join(current).strip())
    return pages


def load_corpus_documents(corpus_dir: Path) -> list[tuple[str, bytes]]:
    """Returns `[(filename, data), ...]` in the fixed order
    `documents.index_document` is called with -- `filename` is what
    `index_document`/`Passage.filename`/`expected_source` all see, never the
    committed source's own name for the DOCX/PDF pair."""
    built: list[tuple[str, bytes]] = []
    for source_name, indexed_filename, kind in _CORPUS_SOURCES:
        source_path = corpus_dir / source_name
        if kind == "text":
            built.append((indexed_filename, source_path.read_bytes()))
        elif kind == "docx":
            text = source_path.read_text(encoding="utf-8-sig")
            built.append((indexed_filename, markdown_to_docx_bytes(text)))
        elif kind == "pdf":
            text = source_path.read_text(encoding="utf-8-sig")
            built.append((indexed_filename, write_pdf(split_pdf_pages(text))))
        else:  # pragma: no cover -- _CORPUS_SOURCES is a fixed, own literal
            raise AssertionError(f"unknown corpus source kind: {kind!r}")
    return built


def load_questions(path: Path) -> list[dict]:
    return json.loads(path.read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# Indexing -- the real pipeline, DOC-05's required `started_at`, a no-op
# progress sink (this script owns no status message).
# ---------------------------------------------------------------------------


def index_corpus(
    conn,
    *,
    embedder,
    documents_to_index: list[tuple[str, bytes]],
    monotonic: Callable[[], float] = time.monotonic,
) -> None:
    for filename, data in documents_to_index:
        documents.index_document(
            conn,
            user_id=EVAL_USER_ID,
            filename=filename,
            data=data,
            embedder=embedder,
            progress=lambda _msg: None,
            now=storage.utc_now_iso(),
            started_at=monotonic(),
            monotonic=monotonic,
        )


# ---------------------------------------------------------------------------
# Retrieval modes (REQ-V190-EVAL-03): `vector` and `hybrid` are id-level --
# hydrated through `storage.chunks_by_ids` exactly once per (question, mode),
# original ranked-id order restored, ids the query didn't return dropped.
# `hybrid+rerank` is `Searcher.search` -- already hydrated `Passage`s.
# ---------------------------------------------------------------------------


def _hydrate_ordered(conn, ids: list[int]) -> list[dict]:
    rows = storage.chunks_by_ids(conn, user_id=EVAL_USER_ID, ids=ids)
    by_id = {row["id"]: row for row in rows}
    return [
        {"filename": by_id[i]["filename"], "page": by_id[i]["page"], "text": by_id[i]["text"]}
        for i in ids
        if i in by_id
    ]


def vector_mode(conn, *, embedder, query: str) -> list[dict]:
    ids = rag.vector_search(conn, user_id=EVAL_USER_ID, embedder=embedder, query=query, k=5)
    return _hydrate_ordered(conn, ids)


def hybrid_mode(conn, *, embedder, query: str) -> list[dict]:
    vector_ids = rag.vector_search(conn, user_id=EVAL_USER_ID, embedder=embedder, query=query)
    rows = storage.user_chunks(conn, user_id=EVAL_USER_ID)
    bm25_ids = rag.bm25_search(rows, query)
    fused = rag.rrf([vector_ids, bm25_ids])[:5]
    return _hydrate_ordered(conn, fused)


def _passages_as_dicts(passages: list[rag.Passage]) -> list[dict]:
    return [{"filename": p.filename, "page": p.page, "text": p.text} for p in passages]


# ---------------------------------------------------------------------------
# Scoring: "a hit is evidence, not a filename."
# ---------------------------------------------------------------------------


def first_hit(
    items: list[dict], *, expected_source: str, expected_evidence: str
) -> tuple[int | None, int | None]:
    """1-based rank and page of the first item whose `filename` equals
    `expected_source` and whose `text` contains `expected_evidence`
    (case-insensitive, whitespace-normalised); `(None, None)` for a miss."""
    for rank, item in enumerate(items, start=1):
        if item["filename"] == expected_source and contains_evidence(
            item["text"], expected_evidence
        ):
            return rank, item["page"]
    return None, None


# ---------------------------------------------------------------------------
# The conversation-aware smoke test (REQ-V190-TOOL-06), advisory only.
# ---------------------------------------------------------------------------


class _RecordingSearcher:
    """Wraps any `.search(query) -> SearchResult` object, recording every
    query it was called with -- this script's only way to inspect what query
    the agent actually issued a turn on, since `tool_calls` rows carry char
    counts, never the query text itself. v1.9.4 T2: also keeps the
    `SearchResult` returned for each call, so turn 3's verdict can check
    whether any call's returned passages included the gold source."""

    def __init__(self, inner) -> None:
        self._inner = inner
        self.queries: list[str] = []
        self.results: list[rag.SearchResult] = []

    def search(self, query: str):
        self.queries.append(query)
        result = self._inner.search(query)
        self.results.append(result)
        return result


def _refusing_runner(  # pragma: no cover -- never invoked
    argv: list[str],  # skylos: ignore -- CommandRunner signature
) -> dict:
    return {"error": "exec is not available in devtools/rag_eval.py"}


def conversation_smoke(
    conn,
    *,
    question: str,
    embedder,
    llm,
    cfg,
    resolve_cost=None,
    run_agent_outcome: Callable = agent.run_agent_outcome,
    rerank_llm=None,
) -> tuple[bool, str, bool, str]:
    """Three live turns: `question` (a real answerable item's text -- the
    caller's job, not a fact this module hardcodes, so the smoke test stays
    corpus-agnostic and runs cleanly against a synthetic corpus offline
    too), then the literal follow-up `«а в неделях?»`, then
    `_SMOKE_FOLLOWUP_2`. Returns `(tool06_ok, tool06_detail, context_proof_ok,
    context_proof_detail)` -- two independent, both-advisory verdicts:

    - TOOL-06 pin: `pass` when the second turn's recorded `search_documents`
      query shares a stemmed token with the first question's text -- evidence
      the follow-up was resolved with the first turn's context in view, not
      asked in isolation. Unchanged from v1.9.3 T2 (see module docstring
      note above `_SMOKE_FOLLOWUP_2` for why this alone can no longer fail
      for a real reason).
    - context-proof: `pass` iff the third turn made at least one
      `search_documents` call **and** at least one of its calls' returned
      passages carries the gold source `_SMOKE_GOLD_SOURCE_2` **and**
      contains the gold evidence `_SMOKE_GOLD_EVIDENCE_2` (`contains_evidence`,
      the same filename-AND-evidence matcher `first_hit` uses to score the
      pinned questions). v1.9.4 review (finding 3): this proves the third
      turn's search actually retrieved the passage that answers the
      transfer question -- both filename and evidence -- not merely that
      some vacation-related page was found (turn 1 already proves that much
      against the same file). It still does NOT prove the query *used*
      turns 1-2's context to get there -- a query that happens to name the
      transfer rule directly, without drawing on any prior-turn context,
      would also pass. The TOOL-06 pin above covers "context was carried
      between turns"; this verdict covers "the agent went back to the
      documents for the new, specific fact."

    v1.9.3 T2 (docs/spec/task-briefs/v193-T2.md): each turn's `Searcher`
    reranks through `rerank_llm` when one is configured, `llm` (the chat/
    agent completion client `_turn` also uses) otherwise -- the same
    `llm=rerank_llm or llm` routing `run()`'s own scored
    `hybrid_rerank_searcher` and `bot.py`'s live searcher already use.
    Before this fix both Searchers here always reranked through the chat
    client regardless -- the LM Studio-primary failover client, not the
    routed rerank model -- which is exactly what put the smoke turn's own
    rerank calls at LM Studio's much slower measured latency instead of
    the fast reranker every other item in this gate uses."""
    conv_id = storage.get_or_create_active_conversation(conn, EVAL_USER_ID)

    def _turn(text: str, searcher) -> None:
        storage.add_user_message(conn, conv_id, text)
        run_agent_outcome(
            conn=conn,
            conv_id=conv_id,
            llm=llm,
            skills={},
            runner=_refusing_runner,
            now=storage.utc_now_iso(),
            cfg=cfg,
            searcher=searcher,
            resolve_cost=resolve_cost,
        )

    searcher1 = _RecordingSearcher(
        rag.Searcher(
            conn,
            user_id=EVAL_USER_ID,
            embedder=embedder,
            llm=rerank_llm or llm,
            cfg=cfg,
            conv_id=conv_id,
            resolve_cost=resolve_cost,
        )
    )
    _turn(question, searcher1)

    searcher2 = _RecordingSearcher(
        rag.Searcher(
            conn,
            user_id=EVAL_USER_ID,
            embedder=embedder,
            llm=rerank_llm or llm,
            cfg=cfg,
            conv_id=conv_id,
            resolve_cost=resolve_cost,
        )
    )
    _turn(_SMOKE_FOLLOWUP, searcher2)

    tool06_ok, tool06_detail = False, ""
    first_tokens = set(rag.tokenize(question))
    for query in searcher2.queries:
        if first_tokens & set(rag.tokenize(query)):
            tool06_ok = True
            tool06_detail = f"turn 2 query {query!r} shares a token with turn 1's question"
            break
    if not tool06_ok:
        tool06_detail = (
            f"turn 2 recorded no search_documents call sharing a token with turn 1's question "
            f"(turn 2 queries: {searcher2.queries!r})"
        )

    searcher3 = _RecordingSearcher(
        rag.Searcher(
            conn,
            user_id=EVAL_USER_ID,
            embedder=embedder,
            llm=rerank_llm or llm,
            cfg=cfg,
            conv_id=conv_id,
            resolve_cost=resolve_cost,
        )
    )
    _turn(_SMOKE_FOLLOWUP_2, searcher3)

    turn3_sources = [[p.filename for p in result.passages] for result in searcher3.results]
    turn3_hit = any(
        p.filename == _SMOKE_GOLD_SOURCE_2 and contains_evidence(p.text, _SMOKE_GOLD_EVIDENCE_2)
        for result in searcher3.results
        for p in result.passages
    )
    if searcher3.queries and turn3_hit:
        context_proof_ok = True
        context_proof_detail = (
            f"turn 3 queries {searcher3.queries!r} returned the gold source "
            f"{_SMOKE_GOLD_SOURCE_2!r} (top sources: {turn3_sources!r})"
        )
    elif not searcher3.queries:
        context_proof_ok = False
        context_proof_detail = "turn 3 recorded no search_documents call"
    else:
        context_proof_ok = False
        context_proof_detail = (
            f"turn 3 queries {searcher3.queries!r} did not return the gold source "
            f"{_SMOKE_GOLD_SOURCE_2!r} (top sources: {turn3_sources!r})"
        )

    return tool06_ok, tool06_detail, context_proof_ok, context_proof_detail


# ---------------------------------------------------------------------------
# The run itself: EVAL-03's core, parameterised over every I/O boundary so
# `tests/test_v190_eval.py` can drive it with `FakeEmbedder`/`FakeLLM`.
# ---------------------------------------------------------------------------


def run(
    *,
    conn,
    cfg: Config,
    embedder,
    llm,
    corpus_dir: Path = CORPUS_DIR,
    questions_path: Path = QUESTIONS_PATH,
    resolve_cost=None,
    run_agent_outcome: Callable = agent.run_agent_outcome,
    print_fn: Callable[[str], None] = print,
    rerank_llm=None,
    eval_chat_llm=None,
) -> int:
    # RET-06/-07 must actually run: force the override regardless of the
    # deployment's own .env, and prove it in the script's own output.
    cfg = dataclasses.replace(cfg, rag_rerank="on", rag_top_k=5)
    print_fn(f"gate-7 override: rag_rerank={cfg.rag_rerank!r} rag_top_k={cfg.rag_top_k!r}")

    hashes = freeze(corpus_dir, questions_path)
    print_fn("frozen sha256:")
    for name, digest in sorted(hashes.items()):
        print_fn(f"  {name}: {digest}")

    questions = load_questions(questions_path)
    answerable = [q for q in questions if q["expected_source"] is not None]
    null_items = [q for q in questions if q["expected_source"] is None]

    try:
        documents_to_index = load_corpus_documents(corpus_dir)
        index_corpus(conn, embedder=embedder, documents_to_index=documents_to_index)
    except (EmbeddingError, EmbeddingTimeoutError, httpx.HTTPError, OSError) as exc:
        print_fn(f"gate-7: FAIL indexing the corpus -- {config.redact(str(exc))}")
        return 2

    ranks: dict[str, list[tuple[int | None, int | None]]] = {mode: [] for mode in MODES}
    rerank_flags: list[tuple[dict, rag.SearchResult]] = []
    hybrid_rerank_searcher = rag.Searcher(
        conn,
        user_id=EVAL_USER_ID,
        embedder=embedder,
        llm=rerank_llm or llm,
        cfg=cfg,
        conv_id=storage.get_or_create_active_conversation(conn, EVAL_USER_ID),
        resolve_cost=resolve_cost,
    )

    try:
        for question in answerable:
            q_text = question["question"]
            vector_items = vector_mode(conn, embedder=embedder, query=q_text)
            hybrid_items = hybrid_mode(conn, embedder=embedder, query=q_text)
            rerank_result = hybrid_rerank_searcher.search(q_text)
            rerank_items = _passages_as_dicts(rerank_result.passages)
            rerank_flags.append((question, rerank_result))

            ranks["vector"].append(
                first_hit(
                    vector_items,
                    expected_source=question["expected_source"],
                    expected_evidence=question["expected_evidence"],
                )
            )
            ranks["hybrid"].append(
                first_hit(
                    hybrid_items,
                    expected_source=question["expected_source"],
                    expected_evidence=question["expected_evidence"],
                )
            )
            ranks["hybrid+rerank"].append(
                first_hit(
                    rerank_items,
                    expected_source=question["expected_source"],
                    expected_evidence=question["expected_evidence"],
                )
            )
    except (EmbeddingError, EmbeddingTimeoutError, LLMError, httpx.HTTPError, OSError) as exc:
        print_fn(f"gate-7: FAIL retrieval -- {config.redact(str(exc))}")
        return 2

    null_advisory: dict[str, list[str]] = {mode: [] for mode in MODES}
    for question in null_items:
        q_text = question["question"]
        vector_items = vector_mode(conn, embedder=embedder, query=q_text)
        hybrid_items = hybrid_mode(conn, embedder=embedder, query=q_text)
        rerank_items = _passages_as_dicts(hybrid_rerank_searcher.search(q_text).passages)
        for mode, items in (
            ("vector", vector_items),
            ("hybrid", hybrid_items),
            ("hybrid+rerank", rerank_items),
        ):
            verdict = "passages returned" if items else "no passage from any file"
            null_advisory[mode].append(f"{q_text!r}: {verdict}")

    # ---- metrics ------------------------------------------------------
    n = len(answerable)
    pdf_indices = [i for i, q in enumerate(answerable) if q["expected_page"] is not None]
    metrics: dict[str, dict[str, float]] = {}
    for mode in MODES:
        mode_ranks = ranks[mode]
        hits = sum(1 for rank, _page in mode_ranks if rank is not None)
        recall = hits / n if n else 0.0
        reciprocal_sum = sum(1.0 / rank if rank is not None else 0.0 for rank, _page in mode_ranks)
        mrr = reciprocal_sum / n if n else 0.0
        page_hits = sum(
            1
            for i in pdf_indices
            if mode_ranks[i][0] is not None and mode_ranks[i][1] == answerable[i]["expected_page"]
        )
        page_hit_rate = page_hits / len(pdf_indices) if pdf_indices else 0.0
        metrics[mode] = {"recall@5": recall, "mrr": mrr, "page_hit_rate": page_hit_rate}

    # ---- table ----------------------------------------------------------
    print_fn("")
    print_fn("| question | " + " | ".join(MODES) + " |")
    print_fn("| --- | " + " | ".join("---" for _ in MODES) + " |")
    for i, question in enumerate(answerable):
        cells = []
        for mode in MODES:
            rank, _page = ranks[mode][i]
            cells.append(f"hit@{rank}" if rank is not None else "miss")
        print_fn(f"| {question['question']} | " + " | ".join(cells) + " |")

    print_fn("")
    print_fn("summary:")
    for mode in MODES:
        m = metrics[mode]
        print_fn(
            f"  {mode}: recall@5={m['recall@5']:.3f} mrr={m['mrr']:.3f} "
            f"page_hit_rate={m['page_hit_rate']:.3f}"
        )

    print_fn("")
    print_fn("null items (advisory, never scored):")
    for mode in MODES:
        for line in null_advisory[mode]:
            print_fn(f"  [{mode}] {line}")

    # ---- conversation-aware smoke (advisory, spends tokens) -------------
    print_fn("")
    print_fn(
        "conversation-aware smoke (advisory): this is the one place in the "
        "whole gate set that spends live inference tokens."
    )
    if answerable:
        # v1.9.4 T2: the smoke turns' chat client is `eval_chat_llm` when
        # `LLM_EVAL_CHAT_MODEL` routes it (opt-in, measurement only), the
        # production `llm` otherwise -- unset, this is exactly `llm`, so the
        # default code path is unchanged.
        tool06_ok, tool06_detail, context_proof_ok, context_proof_detail = conversation_smoke(
            conn,
            question=answerable[0]["question"],
            embedder=embedder,
            llm=eval_chat_llm or llm,
            cfg=cfg,
            resolve_cost=resolve_cost,
            run_agent_outcome=run_agent_outcome,
            rerank_llm=rerank_llm,
        )
        print_fn(
            f"  conversation-aware smoke (TOOL-06 pin): "
            f"{'pass' if tool06_ok else 'fail'} -- {tool06_detail}"
        )
        print_fn(
            f"  conversation-aware smoke (context-proof): "
            f"{'pass' if context_proof_ok else 'fail'} -- {context_proof_detail}"
        )
    else:  # pragma: no cover -- questions.json always carries >= 1 answerable item
        print_fn("  conversation-aware smoke: skipped -- no answerable question to seed it")

    # ---- rerank-actually-ran gate (RET-07) -------------------------------
    failing = [
        (q["question"], r.rerank_attempted, r.rerank_succeeded, r.rerank_failure)
        for q, r in rerank_flags
        if not (r.rerank_attempted and r.rerank_succeeded)
    ]
    if failing:
        print_fn("")
        print_fn("gate-7: FAIL rerank did not run for every answerable item:")
        for q_text, attempted, succeeded, failure in failing:
            print_fn(
                f"  {q_text!r}: rerank_attempted={attempted} rerank_succeeded={succeeded} "
                f"rerank_failure={failure!r}"
            )
        return 2

    if metrics["hybrid"]["recall@5"] < RECALL_FLOOR:
        print_fn("")
        print_fn(
            f"gate-7: FAIL hybrid recall@5={metrics['hybrid']['recall@5']:.3f} "
            f"below the {RECALL_FLOOR} floor"
        )
        return 1

    print_fn("")
    print_fn("gate-7: PASS")
    return 0


# ---------------------------------------------------------------------------
# main(): the real wiring, mirroring `bot.py`'s own construction.
# ---------------------------------------------------------------------------


def main() -> int:
    # v1.9.4 T1: this entry point configured no root logger of its own, so
    # its own warnings went to `logging`'s unredacted `lastResort` stderr
    # handler -- a `basicConfig`-equivalent through the shared helper, same
    # guard `basicConfig` itself uses (act only when root has no handler
    # yet), gives gate 7's own log lines the same redaction as every other
    # entry point.
    root_logger = logging.getLogger()
    if not root_logger.handlers:
        handler = logging.StreamHandler(sys.stderr)
        config.install_redacting_logging(handler, "%(asctime)s %(levelname)s %(name)s %(message)s")
        root_logger.addHandler(handler)
        # v1.9.4 review (finding 1): WARNING, not INFO -- the brief asked
        # only for gate 7's *warnings* to be redacted; INFO buries
        # checks.py's five-line stderr tail on a failing gate under ~88
        # INFO lines/run (httpx POSTs, storage llm_call rows).
        root_logger.setLevel(logging.WARNING)
    try:
        cfg = load_config()
    except config.ConfigError as exc:
        print(f"gate-7: FAIL configuration -- {config.redact(str(exc))}")
        return 2
    if not cfg.rag_enabled:
        print("gate-7: FAIL EMBEDDING_MODEL and EMBEDDING_DIM are not set")
        return 2

    client = httpx.Client()
    try:
        try:
            llm = build_llm_client(cfg, client=client)
            # v1.9.1 T1: mirrors bot.py's own startup wiring -- a bare client
            # on LLM_RERANK_MODEL's provider when configured, so the gate
            # measures the same reranker a live deployment would use.
            rerank_llm = (
                build_llm_client(cfg, client=client, purpose="rerank")
                if cfg.llm_rerank_model
                else None
            )
            # v1.9.4 T2 (docs/spec/task-briefs/v194-T2.md): the opt-in
            # measurement route for the smoke turns' chat completions --
            # unset (the default), this stays None and run() falls through
            # to the production `llm` client exactly as before this change.
            eval_chat_llm = (
                build_llm_client(cfg, client=client, purpose="eval-chat")
                if cfg.llm_eval_chat_model
                else None
            )
        except Exception as exc:  # construction failure is exit 2
            print(f"gate-7: FAIL constructing the chat model -- {config.redact(str(exc))}")
            return 2
        embedder = EmbeddingsClient(
            cfg.embedding_base_url,
            cfg.embedding_model,
            cfg.embedding_dim,
            cfg.embedding_timeout_s,
            client,
            api_key=cfg.embedding_api_key,
        )
        with tempfile.TemporaryDirectory() as tmp_dir:
            db_path = Path(tmp_dir) / "rag_eval.db"
            conn = storage.connect(db_path)
            try:
                storage.init_schema(
                    conn,
                    embedding_dim=cfg.embedding_dim,
                    embedding_model=cfg.embedding_model,
                )
                return run(
                    conn=conn,
                    cfg=cfg,
                    embedder=embedder,
                    llm=llm,
                    rerank_llm=rerank_llm,
                    eval_chat_llm=eval_chat_llm,
                )
            finally:
                conn.close()
    finally:
        client.close()


if __name__ == "__main__":
    raise SystemExit(main())
